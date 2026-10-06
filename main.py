import os
import re
import logging
import sys
import tkinter as tk
from contextlib import closing
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from pathlib import Path
from tkinter import filedialog, messagebox, scrolledtext, simpledialog, ttk
from zipfile import BadZipFile

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
import psycopg
import pyodbc
from psycopg import sql
# import xlwt  # Exportação XLS desativada temporariamente.

# Usa a pasta do executável empacotado ou, durante o desenvolvimento, a pasta deste script.
PASTA_APLICATIVO = (
    Path(sys.executable).resolve().parent
    if getattr(sys, "frozen", False)
    else Path(__file__).resolve().parent
)
ARQUIVO_LOG = PASTA_APLICATIVO / "importador.txt"
ARQUIVO_RELATORIO_VALIDACAO = PASTA_APLICATIVO / "relatorio_validacao_importacao.txt"

LOGGER = logging.getLogger("importador")
LOGGER.setLevel(logging.INFO)
if not LOGGER.handlers:
    manipulador_log = logging.FileHandler(ARQUIVO_LOG, encoding="utf-8")
    formato_log = logging.Formatter(
        "%(asctime)s | %(levelname)s | %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    manipulador_log.setFormatter(formato_log)
    LOGGER.addHandler(manipulador_log)
    LOGGER.propagate = False

# Carrega as variáveis locais do arquivo .env sem versionar credenciais.
load_dotenv()


# Reúne a janela principal, os menus e as operações com o PostgreSQL.
class ImportadorApp:
    def __init__(self, root):
        # Prepara a janela, o estado da conexão e os elementos da tela inicial.
        self.root = root
        self.root.title("Importador")
        self.root.geometry("600x340")
        self.root.minsize(460, 260)
        self.configuracao = None
        self.configuracao_sqlserver = None
        self.status = tk.StringVar(value="Configure os dados do PostgreSQL para começar.")
        self.root.report_callback_exception = self._registrar_erro_interface

        self._criar_menu()

        conteudo = ttk.Frame(root, padding=24)
        conteudo.pack(fill="both", expand=True)
        ttk.Label(conteudo, text="Importador", font=("Segoe UI", 18, "bold")).pack(anchor="w")
        ttk.Label(
            conteudo,
            text="Use o menu Obter dados para consultar tabelas e salvar em Excel.",
        ).pack(anchor="w", pady=(8, 20))
        ttk.Separator(conteudo).pack(fill="x", pady=(0, 14))
        ttk.Label(conteudo, textvariable=self.status, wraplength=520).pack(anchor="w")
        LOGGER.info("Aplicativo iniciado.")

    @staticmethod
    def _registrar_erro_interface(tipo_erro, _valor_erro, _traceback_erro):
        LOGGER.error("Erro não tratado na interface: %s", tipo_erro.__name__)

    def _criar_menu(self):
        # Conecta cada opção de menu ao fluxo correspondente do aplicativo.
        menu = tk.Menu(self.root)
        menu_geral = tk.Menu(menu, tearoff=False)
        menu_geral.add_command(label="Configurar conexão...", command=self.abrir_configuracao)
        menu_geral.add_command(label="Testar conexão", command=self.testar_conexao_salva)
        menu_geral.add_separator()
        menu_geral.add_command(
            label="Executar comandos...", command=self.abrir_execucao_comandos
        )
        menu.add_cascade(label="Menu", menu=menu_geral)

        obter_dados_menu = tk.Menu(menu, tearoff=False)
        postgres_menu = tk.Menu(obter_dados_menu, tearoff=False)
        postgres_menu.add_command(
            label="Consultar tabela e salvar em Excel...",
            command=self.abrir_obter_dados,
        )
        obter_dados_menu.add_cascade(label="PostgreSQL", menu=postgres_menu)

        sql_server_menu = tk.Menu(obter_dados_menu, tearoff=False)
        sql_server_menu.add_command(
            label="Consultar tabela e salvar em Excel...",
            command=self.abrir_obter_dados_sqlserver,
        )
        obter_dados_menu.add_cascade(label="SQL Server", menu=sql_server_menu)
        menu.add_cascade(label="Obter dados", menu=obter_dados_menu)

        importar_menu = tk.Menu(menu, tearoff=False)
        importar_menu.add_command(
            label="Importar dados por XLSX...", command=self.importar_dados_xlsx
        )
        menu.add_cascade(label="Importar", menu=importar_menu)

        ajuda_menu = tk.Menu(menu, tearoff=False)
        ajuda_menu.add_command(
            label="Sobre",
            command=lambda: messagebox.showinfo(
                "Sobre",
                "Importador de dados para PostgreSQL e SQL Server.\nVersão Beta 1.3",
                parent=self.root,
            ),
        )
        menu.add_cascade(label="Ajuda", menu=ajuda_menu)
        self.root.config(menu=menu)

    def abrir_obter_dados(self):
        # Monta a tela para consultar uma tabela e exportar o resultado para Excel.
        janela = tk.Toplevel(self.root)
        janela.title("Obter dados do PostgreSQL")
        janela.geometry("960x640")
        janela.minsize(700, 480)
        janela.transient(self.root)

        # Reaproveita a conexão da sessão ou preenche os campos com valores do .env.
        valores_atuais = self.configuracao or {
            "host": os.getenv("PGHOST", "localhost"),
            "port": os.getenv("PGPORT", "5432"),
            "dbname": os.getenv("PGDATABASE", "autosystem"),
            "user": os.getenv("PGUSER", "postgres"),
            "password": os.getenv("PGPASSWORD", ""),
        }
        campos = (
            ("Servidor", "host"),
            ("Porta", "port"),
            ("Banco", "dbname"),
            ("Usuário", "user"),
            ("Senha", "password"),
        )
        variaveis = {
            chave: tk.StringVar(value=str(valores_atuais[chave]))
            for _, chave in campos
        }
        nome_tabela = tk.StringVar()
        status = tk.StringVar(value="Informe a conexão e a tabela para consultar.")
        resultado = {"colunas": [], "linhas": []}

        # Cria os campos de conexão e a entrada do nome da tabela.
        formulario = ttk.Frame(janela, padding=12)
        formulario.pack(fill="x")
        for indice, (rotulo, chave) in enumerate(campos):
            coluna = indice * 2
            ttk.Label(formulario, text=rotulo).grid(
                row=0, column=coluna, sticky="w", padx=(0, 6)
            )
            ttk.Entry(
                formulario,
                textvariable=variaveis[chave],
                width=16,
                show="*" if chave == "password" else "",
            ).grid(row=1, column=coluna, sticky="ew", padx=(0, 10), pady=(4, 0))
            formulario.columnconfigure(coluna, weight=1)

        ttk.Label(formulario, text="Tabela (ou esquema.tabela)").grid(
            row=2, column=0, columnspan=2, sticky="w", pady=(12, 0)
        )
        ttk.Entry(formulario, textvariable=nome_tabela, width=36).grid(
            row=3, column=0, columnspan=2, sticky="ew", padx=(0, 10), pady=(4, 0)
        )

        # Liga os botões de consulta e exportação às funções de processamento.
        barra = ttk.Frame(janela, padding=(12, 0, 12, 8))
        barra.pack(fill="x")
        botao_excel = ttk.Button(
            barra,
            text="Salvar resultado em Excel...",
            state="disabled",
            command=lambda: self.salvar_dados_excel(resultado, janela),
        )
        botao_excel.pack(side="right")
        ttk.Button(
            barra,
            text="Obter dados",
            command=lambda: self.consultar_tabela(
                variaveis,
                nome_tabela,
                resultado,
                tabela_resultados,
                status,
                botao_excel,
                janela,
            ),
        ).pack(side="left")
        ttk.Label(barra, textvariable=status).pack(side="left", padx=12)

        # Exibe o resultado em uma grade com rolagem horizontal e vertical.
        area_resultados = ttk.Frame(janela, padding=(12, 0, 12, 12))
        area_resultados.pack(fill="both", expand=True)
        tabela_resultados = ttk.Treeview(area_resultados, show="headings")
        rolagem_vertical = ttk.Scrollbar(
            area_resultados, orient="vertical", command=tabela_resultados.yview
        )
        rolagem_horizontal = ttk.Scrollbar(
            area_resultados, orient="horizontal", command=tabela_resultados.xview
        )
        tabela_resultados.configure(
            yscrollcommand=rolagem_vertical.set,
            xscrollcommand=rolagem_horizontal.set,
        )
        tabela_resultados.grid(row=0, column=0, sticky="nsew")
        rolagem_vertical.grid(row=0, column=1, sticky="ns")
        rolagem_horizontal.grid(row=1, column=0, sticky="ew")
        area_resultados.rowconfigure(0, weight=1)
        area_resultados.columnconfigure(0, weight=1)

    def abrir_obter_dados_sqlserver(self):
        # Abre a tela de consulta usando SQL Server e os drivers ODBC instalados.
        drivers = pyodbc.drivers()
        driver_preferido = os.getenv("MSSQL_DRIVER", "").strip()
        if not driver_preferido:
            driver_preferido = next(
                (
                    driver
                    for versao in ("ODBC Driver 18 for SQL Server", "ODBC Driver 17 for SQL Server")
                    for driver in drivers
                    if driver == versao
                ),
                drivers[0] if drivers else "ODBC Driver 18 for SQL Server",
            )
        if driver_preferido not in drivers:
            drivers.insert(0, driver_preferido)

        janela = tk.Toplevel(self.root)
        janela.title("Obter dados do SQL Server")
        janela.geometry("960x640")
        janela.minsize(700, 480)
        janela.transient(self.root)

        valores_atuais = self.configuracao_sqlserver or {
            "driver": driver_preferido,
            "server": os.getenv("MSSQL_SERVER", "localhost"),
            "port": os.getenv("MSSQL_PORT", "1433"),
            "dbname": os.getenv("MSSQL_DATABASE", ""),
            "user": os.getenv("MSSQL_USER", ""),
            "password": os.getenv("MSSQL_PASSWORD", ""),
            "trust_server_certificate": os.getenv(
                "MSSQL_TRUST_SERVER_CERTIFICATE", "false"
            ).lower() in {"1", "true", "yes"},
        }
        campos = (
            ("Driver ODBC", "driver"),
            ("Servidor", "server"),
            ("Porta", "port"),
            ("Banco", "dbname"),
            ("Usuário", "user"),
            ("Senha", "password"),
        )
        variaveis = {
            chave: tk.StringVar(value=str(valores_atuais[chave]))
            for _, chave in campos
        }
        confiar_certificado = tk.BooleanVar(
            value=valores_atuais["trust_server_certificate"]
        )
        nome_tabela = tk.StringVar()
        status = tk.StringVar(value="Informe a conexão e a tabela para consultar.")
        resultado = {"colunas": [], "linhas": []}

        formulario = ttk.Frame(janela, padding=12)
        formulario.pack(fill="x")
        for indice, (rotulo, chave) in enumerate(campos):
            coluna = indice * 2
            ttk.Label(formulario, text=rotulo).grid(
                row=0, column=coluna, sticky="w", padx=(0, 6)
            )
            if chave == "driver":
                entrada = ttk.Combobox(
                    formulario,
                    textvariable=variaveis[chave],
                    values=drivers,
                    width=24,
                    state="readonly",
                )
            else:
                entrada = ttk.Entry(
                    formulario,
                    textvariable=variaveis[chave],
                    width=16,
                    show="*" if chave == "password" else "",
                )
            entrada.grid(row=1, column=coluna, sticky="ew", padx=(0, 10), pady=(4, 0))
            formulario.columnconfigure(coluna, weight=1)

        ttk.Checkbutton(
            formulario,
            text="Confiar no certificado do servidor (desenvolvimento)",
            variable=confiar_certificado,
        ).grid(row=2, column=0, columnspan=6, sticky="w", pady=(10, 0))
        ttk.Label(formulario, text="Tabela (ou esquema.tabela)").grid(
            row=3, column=0, columnspan=2, sticky="w", pady=(10, 0)
        )
        ttk.Entry(formulario, textvariable=nome_tabela, width=36).grid(
            row=4, column=0, columnspan=2, sticky="ew", padx=(0, 10), pady=(4, 0)
        )

        barra = ttk.Frame(janela, padding=(12, 0, 12, 8))
        barra.pack(fill="x")
        botao_excel = ttk.Button(
            barra,
            text="Salvar resultado em Excel...",
            state="disabled",
            command=lambda: self.salvar_dados_excel(resultado, janela),
        )
        botao_excel.pack(side="right")
        ttk.Button(
            barra,
            text="Obter dados",
            command=lambda: self.consultar_tabela_sqlserver(
                variaveis,
                confiar_certificado,
                nome_tabela,
                resultado,
                tabela_resultados,
                status,
                botao_excel,
                janela,
            ),
        ).pack(side="left")
        ttk.Label(barra, textvariable=status).pack(side="left", padx=12)

        area_resultados = ttk.Frame(janela, padding=(12, 0, 12, 12))
        area_resultados.pack(fill="both", expand=True)
        tabela_resultados = ttk.Treeview(area_resultados, show="headings")
        rolagem_vertical = ttk.Scrollbar(
            area_resultados, orient="vertical", command=tabela_resultados.yview
        )
        rolagem_horizontal = ttk.Scrollbar(
            area_resultados, orient="horizontal", command=tabela_resultados.xview
        )
        tabela_resultados.configure(
            yscrollcommand=rolagem_vertical.set,
            xscrollcommand=rolagem_horizontal.set,
        )
        tabela_resultados.grid(row=0, column=0, sticky="nsew")
        rolagem_vertical.grid(row=0, column=1, sticky="ns")
        rolagem_horizontal.grid(row=1, column=0, sticky="ew")
        area_resultados.rowconfigure(0, weight=1)
        area_resultados.columnconfigure(0, weight=1)

    @staticmethod
    def _valor_conexao_odbc(valor):
        # Escapa valores de conexão ODBC que contenham chaves ou ponto e vírgula.
        return "{" + str(valor).replace("}", "}}") + "}"

    @staticmethod
    def _identificador_sqlserver(partes):
        # Delimita cada parte do nome para evitar interpretar o nome como SQL executável.
        return ".".join("[" + parte.replace("]", "]]" ) + "]" for parte in partes)

    def consultar_tabela_sqlserver(
        self,
        variaveis,
        confiar_certificado,
        nome_tabela,
        resultado,
        tabela_resultados,
        status,
        botao_excel,
        parent,
    ):
        # Valida os dados ODBC e o destino antes de abrir a conexão SQL Server.
        configuracao = {chave: variavel.get().strip() for chave, variavel in variaveis.items()}
        try:
            configuracao["port"] = int(configuracao["port"])
            if not 1 <= configuracao["port"] <= 65535:
                raise ValueError
        except ValueError:
            messagebox.showerror("Porta inválida", "Informe uma porta entre 1 e 65535.", parent=parent)
            return

        obrigatorios = ("driver", "server", "dbname", "user", "password")
        if any(not configuracao[chave] for chave in obrigatorios):
            messagebox.showerror(
                "Dados incompletos",
                "Preencha driver, servidor, banco, usuário e senha.",
                parent=parent,
            )
            return

        partes_tabela = [parte.strip() for parte in nome_tabela.get().strip().split(".")]
        if not partes_tabela[0] or len(partes_tabela) > 2 or any(not parte for parte in partes_tabela):
            messagebox.showerror(
                "Tabela inválida",
                "Informe o nome da tabela ou esquema.tabela.",
                parent=parent,
            )
            return

        configuracao["trust_server_certificate"] = confiar_certificado.get()
        nome_tabela_log = ".".join(partes_tabela)
        LOGGER.info(
            "Consulta SQL Server iniciada: tabela=%s banco=%s servidor=%s.",
            nome_tabela_log,
            configuracao["dbname"],
            configuracao["server"],
        )
        tabela_resultados.delete(*tabela_resultados.get_children())
        resultado["colunas"] = []
        resultado["linhas"] = []
        botao_excel.state(["disabled"])

        servidor = f"{configuracao['server']},{configuracao['port']}"
        string_conexao = ";".join(
            (
                f"DRIVER={self._valor_conexao_odbc(configuracao['driver'])}",
                f"SERVER={self._valor_conexao_odbc(servidor)}",
                f"DATABASE={self._valor_conexao_odbc(configuracao['dbname'])}",
                f"UID={self._valor_conexao_odbc(configuracao['user'])}",
                f"PWD={self._valor_conexao_odbc(configuracao['password'])}",
                "Encrypt=yes",
                "TrustServerCertificate=" + ("yes" if confiar_certificado.get() else "no"),
                "Connection Timeout=5",
            )
        )
        tabela_sql = self._identificador_sqlserver(partes_tabela)
        consulta = f"SELECT * FROM {tabela_sql}"

        try:
            with closing(pyodbc.connect(string_conexao, timeout=5)) as conexao:
                with closing(conexao.cursor()) as cursor:
                    cursor.execute(consulta)
                    colunas = [coluna[0] for coluna in cursor.description]
                    linhas = []
                    while True:
                        lote = cursor.fetchmany(2000)
                        if not lote:
                            break
                        linhas.extend(lote)
        except pyodbc.Error as erro:
            status.set("Falha ao obter dados do SQL Server.")
            LOGGER.error(
                "Consulta SQL Server falhou: tabela=%s erro=%s.",
                nome_tabela_log,
                type(erro).__name__,
            )
            messagebox.showerror("Erro na consulta SQL Server", str(erro), parent=parent)
            return

        LOGGER.info(
            "Consulta SQL Server concluída: tabela=%s registros=%s.",
            nome_tabela_log,
            len(linhas),
        )
        self.configuracao_sqlserver = configuracao
        resultado["colunas"] = colunas
        resultado["linhas"] = linhas
        identificadores = [f"coluna_{indice}" for indice in range(len(colunas))]
        tabela_resultados.configure(columns=identificadores)
        for identificador, nome in zip(identificadores, colunas):
            tabela_resultados.heading(identificador, text=nome)
            tabela_resultados.column(identificador, width=150, minwidth=80, stretch=True)
        for linha in linhas[:1000]:
            tabela_resultados.insert(
                "",
                "end",
                values=[
                    "" if valor is None else str(self._decodificar_bytes(valor))
                    for valor in linha
                ],
            )

        botao_excel.state(["!disabled"])
        status.set(
            f"{len(linhas)} registro(s) obtido(s); a tela exibe até 1.000."
        )

    def consultar_tabela(
        self,
        variaveis,
        nome_tabela,
        resultado,
        tabela_resultados,
        status,
        botao_excel,
        parent,
    ):
        # Lê e valida os dados de conexão antes de abrir qualquer conexão.
        configuracao = {chave: variavel.get().strip() for chave, variavel in variaveis.items()}
        try:
            configuracao["port"] = int(configuracao["port"])
            if not 1 <= configuracao["port"] <= 65535:
                raise ValueError
        except ValueError:
            messagebox.showerror("Porta inválida", "Informe uma porta entre 1 e 65535.", parent=parent)
            return

        if any(not configuracao[chave] for chave in ("host", "dbname", "user", "password")):
            messagebox.showerror(
                "Dados incompletos",
                "Preencha servidor, banco, usuário e senha.",
                parent=parent,
            )
            return

        # Aceita tabela ou esquema.tabela; os identificadores serão escapados pelo psycopg.
        partes_tabela = [parte.strip() for parte in nome_tabela.get().strip().split(".")]
        if not partes_tabela[0] or len(partes_tabela) > 2 or any(not parte for parte in partes_tabela):
            messagebox.showerror(
                "Tabela inválida",
                "Informe o nome da tabela ou esquema.tabela.",
                parent=parent,
            )
            return

        nome_tabela_log = ".".join(partes_tabela)
        LOGGER.info(
            "Consulta iniciada: tabela=%s banco=%s servidor=%s.",
            nome_tabela_log,
            configuracao["dbname"],
            configuracao["host"],
        )
        tabela_resultados.delete(*tabela_resultados.get_children())
        resultado["colunas"] = []
        resultado["linhas"] = []
        botao_excel.state(["disabled"])

        # Executa uma consulta somente leitura e busca registros em lotes.
        consulta = sql.SQL("SELECT * FROM {}").format(sql.Identifier(*partes_tabela))
        try:
            with psycopg.connect(**configuracao, connect_timeout=5) as conexao:
                with conexao.cursor() as cursor:
                    cursor.execute("SET TRANSACTION READ ONLY")
                    cursor.execute(consulta)
                    colunas = [coluna.name for coluna in cursor.description]
                    linhas = []
                    while True:
                        lote = cursor.fetchmany(2000)
                        if not lote:
                            break
                        linhas.extend(lote)
        except psycopg.Error as erro:
            status.set("Falha ao obter dados.")
            LOGGER.error(
                "Consulta falhou para tabela=%s (%s).",
                nome_tabela_log,
                type(erro).__name__,
            )
            messagebox.showerror("Erro na consulta", str(erro), parent=parent)
            return

        # Guarda os dados retornados, exibe até 1.000 linhas e habilita a exportação.
        LOGGER.info("Consulta concluída: tabela=%s registros=%s.", nome_tabela_log, len(linhas))
        self.configuracao = configuracao
        resultado["colunas"] = colunas
        resultado["linhas"] = linhas
        identificadores = [f"coluna_{indice}" for indice in range(len(colunas))]
        tabela_resultados.configure(columns=identificadores)
        for identificador, nome in zip(identificadores, colunas):
            tabela_resultados.heading(identificador, text=nome)
            tabela_resultados.column(identificador, width=150, minwidth=80, stretch=True)
        for linha in linhas[:1000]:
            tabela_resultados.insert(
                "",
                "end",
                values=[
                    "" if valor is None else str(self._decodificar_bytes(valor))
                    for valor in linha
                ],
            )

        botao_excel.state(["!disabled"])
        status.set(
            f"{len(linhas)} registro(s) obtido(s); a tela exibe até 1.000."
        )

    def salvar_dados_excel(self, resultado, parent):
        # Verifica se há dados e se o resultado cabe em uma planilha do Excel.
        if not resultado["colunas"]:
            LOGGER.warning("Exportação Excel solicitada sem dados de consulta.")
            messagebox.showwarning("Sem dados", "Faça uma consulta antes de salvar.", parent=parent)
            return
        if len(resultado["linhas"]) > 1_048_575:
            messagebox.showerror(
                "Limite do Excel",
                "O resultado ultrapassa o limite de 1.048.575 registros por aba.",
                parent=parent,
            )
            return

        caminho = filedialog.asksaveasfilename(
            parent=parent,
            title="Salvar resultado da consulta",
            initialfile="dados_consulta.xlsx",
            defaultextension=".xlsx",
            filetypes=[("Planilha Excel", "*.xlsx")],
        )
        if not caminho:
            LOGGER.info("Salvamento Excel cancelado pelo usuário.")
            return

        # Grava cabeçalho e todas as linhas em modo streaming para reduzir memória.
        try:
            workbook = Workbook(write_only=True)
            planilha = workbook.create_sheet("Dados")
            planilha.append(resultado["colunas"])
            tipos_excel = (str, int, float, bool, Decimal, date, datetime, time, timedelta)
            for linha in resultado["linhas"]:
                planilha.append(
                    [
                        self._decodificar_bytes(valor)
                        if valor is None or isinstance(valor, tipos_excel)
                        else str(self._decodificar_bytes(valor))
                        for valor in linha
                    ]
                )
            workbook.save(caminho)
        except Exception as erro:
            LOGGER.error("Falha ao salvar planilha Excel (%s).", type(erro).__name__)
            messagebox.showerror("Erro ao salvar Excel", str(erro), parent=parent)
            return

        LOGGER.info(
            "Planilha Excel salva: arquivo=%s registros=%s.",
            Path(caminho).name,
            len(resultado["linhas"]),
        )
        messagebox.showinfo(
            "Arquivo salvo",
            f"{len(resultado['linhas'])} registro(s) salvo(s) em:\n{caminho}",
            parent=parent,
        )

    @staticmethod
    def _decodificar_bytes(valor):
        # Converte conteúdo binário UTF-8 para texto legível nas grades e planilhas.
        if isinstance(valor, (bytes, bytearray, memoryview)):
            return bytes(valor).decode("utf-8", errors="replace")
        return valor

    @staticmethod
    def _normalizar_valor_unico(valor):
        # Normaliza os tipos numéricos do Excel e ignora espaços/capitalização em textos.
        if isinstance(valor, str):
            return "texto", valor.strip().casefold()
        if isinstance(valor, bool):
            return "booleano", valor
        if isinstance(valor, (int, float, Decimal)):
            return "numero", Decimal(str(valor)).normalize()
        return type(valor).__name__, str(valor).strip().casefold()

    def _analisar_colunas_unicas_xlsx(self, caminho, colunas):
        # Exige codigo preenchido; codigo_Barra pode ficar vazio, mas não pode se repetir preenchido.
        colunas_obrigatorias = ("codigo", "codigo_Barra")
        indices_colunas = {nome.casefold(): [] for nome in colunas_obrigatorias}
        problemas = []

        for indice, nome_coluna in enumerate(colunas):
            chave = nome_coluna.casefold()
            if chave in indices_colunas:
                indices_colunas[chave].append(indice)

        for nome_coluna in colunas_obrigatorias:
            indices = indices_colunas[nome_coluna.casefold()]
            if not indices:
                problemas.append(f'Coluna obrigatória "{nome_coluna}" ausente no cabeçalho.')
            elif len(indices) > 1:
                problemas.append(f'Coluna obrigatória "{nome_coluna}" aparece mais de uma vez.')

        primeira_ocorrencia = {nome: {} for nome in colunas_obrigatorias}
        linhas_analisadas = 0
        workbook = load_workbook(caminho, read_only=True, data_only=True)
        try:
            planilha = workbook.active
            for numero_linha, linha in enumerate(
                planilha.iter_rows(min_row=2, values_only=True), start=2
            ):
                valores = list(linha)
                if not any(valor is not None for valor in valores):
                    continue
                linhas_analisadas += 1

                for nome_coluna in colunas_obrigatorias:
                    indices = indices_colunas[nome_coluna.casefold()]
                    if len(indices) != 1:
                        continue

                    indice = indices[0]
                    valor = valores[indice] if indice < len(valores) else None
                    if valor is None or (isinstance(valor, str) and not valor.strip()):
                        if nome_coluna.casefold() == "codigo_barra":
                            continue
                        problemas.append(f'Linha {numero_linha}, coluna "{nome_coluna}": valor vazio.')
                        continue

                    chave_valor = self._normalizar_valor_unico(valor)
                    ocorrencias = primeira_ocorrencia[nome_coluna]
                    if chave_valor in ocorrencias:
                        primeira_linha, valor_original = ocorrencias[chave_valor]
                        problemas.append(
                            f'Linha {numero_linha}, coluna "{nome_coluna}": valor '
                            f'"{valor_original}" duplicado (primeira ocorrência na linha '
                            f'{primeira_linha}).'
                        )
                    else:
                        ocorrencias[chave_valor] = (numero_linha, valor)
        finally:
            workbook.close()

        return planilha.title, linhas_analisadas, problemas

    def importar_dados_xlsx(self):
        # Solicita destino e arquivo antes de começar a preparar os registros.
        if not self._exigir_configuracao():
            return

        nome_tabela = simpledialog.askstring(
            "Tabela de destino",
            "Informe a tabela onde os dados serão inseridos (ou esquema.tabela):",
            initialvalue="produto",
            parent=self.root,
        )
        if nome_tabela is None:
            LOGGER.info("Importação XLSX cancelada antes da escolha da tabela.")
            return

        # Separa o nome da tabela para montar um identificador SQL seguro.
        partes_tabela = [parte.strip() for parte in nome_tabela.split(".")]
        if not partes_tabela[0] or len(partes_tabela) > 2 or any(not parte for parte in partes_tabela):
            messagebox.showerror(
                "Tabela inválida",
                "Informe o nome da tabela ou esquema.tabela.",
                parent=self.root,
            )
            return
        nome_tabela = ".".join(partes_tabela)
        LOGGER.info("Importação XLSX iniciada: tabela=%s.", nome_tabela)

        caminho = filedialog.askopenfilename(
            parent=self.root,
            title="Selecionar planilha XLSX para importar",
            filetypes=[("Planilhas Excel", "*.xlsx"), ("Todos os arquivos", "*.*")],
        )
        if not caminho:
            LOGGER.info("Seleção de planilha XLSX cancelada.")
            return
        LOGGER.info("Planilha selecionada: arquivo=%s.", Path(caminho).name)

        # Abre a planilha ativa e usa a primeira linha como nomes das colunas.
        try:
            workbook = load_workbook(caminho, read_only=True, data_only=True)
            try:
                planilha = workbook.active
                cabecalho = next(planilha.iter_rows(values_only=True), None)
            finally:
                workbook.close()
        except (OSError, BadZipFile, InvalidFileException, ValueError) as erro:
            LOGGER.error("Falha ao ler planilha XLSX (%s).", type(erro).__name__)
            messagebox.showerror("Erro ao ler XLSX", str(erro), parent=self.root)
            return

        if not cabecalho:
            LOGGER.warning("Importação XLSX cancelada: cabeçalho vazio.")
            messagebox.showerror(
                "Planilha vazia",
                "A aba ativa não contém uma linha de cabeçalho.",
                parent=self.root,
            )
            return

        # Valida o cabeçalho antes de pedir confirmação para inserir os dados.
        colunas = [str(coluna).strip() if coluna is not None else "" for coluna in cabecalho]
        if any(not coluna for coluna in colunas):
            LOGGER.warning("Importação XLSX cancelada: há coluna sem nome.")
            messagebox.showerror(
                "Cabeçalho inválido",
                "Todas as células do cabeçalho precisam ter um nome de coluna.",
                parent=self.root,
            )
            return
        if len(set(colunas)) != len(colunas):
            LOGGER.warning("Importação XLSX cancelada: há nomes de coluna duplicados.")
            messagebox.showerror(
                "Cabeçalho inválido",
                "O cabeçalho contém nomes de coluna duplicados.",
                parent=self.root,
            )
            return

        # Interrompe a importação se codigo falhar ou se codigo_Barra tiver duplicidades.
        try:
            nome_aba, linhas_analisadas, problemas = self._analisar_colunas_unicas_xlsx(
                caminho, colunas
            )
        except (OSError, BadZipFile, InvalidFileException, ValueError) as erro:
            LOGGER.error("Falha na análise prévia da planilha (%s).", type(erro).__name__)
            messagebox.showerror("Erro na análise da planilha", str(erro), parent=self.root)
            return

        if problemas:
            conteudo_relatorio = [
                "RELATORIO DE VALIDACAO DA IMPORTACAO XLSX",
                f"Data e hora: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                f"Arquivo: {Path(caminho).name}",
                f"Aba analisada: {nome_aba}",
                f"Tabela de destino: {nome_tabela}",
                f"Linhas com dados analisadas: {linhas_analisadas}",
                f"Problemas encontrados: {len(problemas)}",
                "",
                "DETALHES",
                *[f"- {problema}" for problema in problemas],
            ]
            try:
                ARQUIVO_RELATORIO_VALIDACAO.write_text(
                    "\n".join(conteudo_relatorio) + "\n",
                    encoding="utf-8",
                )
            except OSError as erro:
                LOGGER.error("Falha ao salvar relatório de validação (%s).", type(erro).__name__)
                messagebox.showerror(
                    "Erro ao salvar relatório",
                    f"Não foi possível salvar o relatório:\n{erro}",
                    parent=self.root,
                )
                return

            LOGGER.warning(
                "Importação bloqueada pela validação: tabela=%s problemas=%s relatório=%s.",
                nome_tabela,
                len(problemas),
                ARQUIVO_RELATORIO_VALIDACAO.name,
            )
            messagebox.showerror(
                "Planilha reprovada na validação",
                f"Foram encontrados {len(problemas)} problema(s) nas colunas "
                "codigo e codigo_Barra preenchido. A importação foi cancelada.\n\n"
                f"Relatório salvo em:\n{ARQUIVO_RELATORIO_VALIDACAO}",
                parent=self.root,
            )
            return

        LOGGER.info(
            "Validação de codigo e codigo_Barra preenchido concluída: tabela=%s linhas=%s.",
            nome_tabela,
            linhas_analisadas,
        )

        nomes_colunas = ", ".join(colunas)
        confirmar = messagebox.askyesno(
            "Confirmar importação",
            f"Colunas lidas da aba ativa:\n{nomes_colunas}\n\n"
            f"Inserir os registros na tabela '{nome_tabela}'?",
            parent=self.root,
        )
        if not confirmar:
            LOGGER.info("Importação XLSX cancelada na confirmação: tabela=%s.", nome_tabela)
            return

        # Monta o INSERT com identificadores escapados e valores parametrizados.
        consulta_insert = sql.SQL("INSERT INTO {} ({}) VALUES ({})").format(
            sql.Identifier(*partes_tabela),
            sql.SQL(", ").join(sql.Identifier(coluna) for coluna in colunas),
            sql.SQL(", ").join(sql.Placeholder() for _ in colunas),
        )
        linhas_importadas = 0

        def ler_linhas():
            # Lê as linhas da planilha sob demanda e valida a quantidade de colunas.
            nonlocal linhas_importadas
            workbook = load_workbook(caminho, read_only=True, data_only=True)
            try:
                planilha = workbook.active
                for numero_linha, linha in enumerate(
                    planilha.iter_rows(min_row=2, values_only=True), start=2
                ):
                    valores = list(linha)
                    if all(valor is None for valor in valores):
                        continue
                    if len(valores) > len(colunas) and any(
                        valor is not None for valor in valores[len(colunas):]
                    ):
                        raise ValueError(
                            f"A linha {numero_linha} contém valores além das colunas do cabeçalho."
                        )
                    valores = valores[:len(colunas)]
                    valores.extend([None] * (len(colunas) - len(valores)))
                    linhas_importadas += 1
                    yield valores
            finally:
                workbook.close()

        # Executa todos os INSERTs na mesma transação para permitir rollback em erro.
        try:
            with psycopg.connect(**self.configuracao, connect_timeout=5) as conexao:
                with conexao.cursor() as cursor:
                    cursor.executemany(consulta_insert, ler_linhas())
        except (psycopg.Error, OSError, BadZipFile, InvalidFileException, ValueError) as erro:
            self.status.set("Falha ao importar o XLSX.")
            LOGGER.error(
                "Importação XLSX falhou: tabela=%s erro=%s.",
                nome_tabela,
                type(erro).__name__,
            )
            messagebox.showerror("Erro na importação", str(erro), parent=self.root)
            return

        LOGGER.info(
            "Importação XLSX concluída: tabela=%s registros=%s.",
            nome_tabela,
            linhas_importadas,
        )
        self.status.set(
            f"Importação concluída: {linhas_importadas} registro(s) inserido(s) em {nome_tabela}."
        )
        messagebox.showinfo(
            "Importação concluída",
            f"{linhas_importadas} registro(s) inserido(s) na tabela '{nome_tabela}'.",
            parent=self.root,
        )

    def abrir_configuracao(self):
        # Abre o formulário modal usado para informar ou testar o acesso ao banco.
        janela = tk.Toplevel(self.root)
        janela.title("Configurar conexão PostgreSQL")
        janela.resizable(False, False)
        janela.transient(self.root)
        janela.grab_set()

        campos = (
            ("Servidor", "host"),
            ("Porta", "port"),
            ("Banco de dados", "dbname"),
            ("Usuário", "user"),
            ("Senha", "password"),
        )
        # Preenche o formulário com os dados da sessão ou com as variáveis locais.
        valores_atuais = self.configuracao or {
            "host": os.getenv("PGHOST", "localhost"),
            "port": os.getenv("PGPORT", "5432"),
            "dbname": os.getenv("PGDATABASE", "autosystem"),
            "user": os.getenv("PGUSER", "postgres"),
            "password": os.getenv("PGPASSWORD", ""),
        }
        variaveis = {
            chave: tk.StringVar(value=str(valores_atuais[chave]))
            for _, chave in campos
        }

        formulario = ttk.Frame(janela, padding=16)
        formulario.grid(row=0, column=0, sticky="nsew")
        for linha, (rotulo, chave) in enumerate(campos):
            ttk.Label(formulario, text=rotulo).grid(
                row=linha, column=0, sticky="w", padx=(0, 12), pady=5
            )
            entrada = ttk.Entry(
                formulario,
                textvariable=variaveis[chave],
                width=34,
                show="*" if chave == "password" else "",
            )
            entrada.grid(row=linha, column=1, sticky="ew", pady=5)

        def ler_configuracao():
            # Valida porta e campos obrigatórios e devolve os dados prontos para uso.
            valores = {chave: variavel.get().strip() for chave, variavel in variaveis.items()}
            try:
                valores["port"] = int(valores["port"])
                if not 1 <= valores["port"] <= 65535:
                    raise ValueError
            except ValueError:
                messagebox.showerror("Porta inválida", "Informe uma porta entre 1 e 65535.", parent=janela)
                return None

            if (
                not valores["host"]
                or not valores["dbname"]
                or not valores["user"]
                or not valores["password"]
            ):
                messagebox.showerror(
                    "Dados incompletos",
                    "Preencha servidor, banco de dados, usuário e senha.",
                    parent=janela,
                )
                return None
            return valores

        def testar():
            # Testa os dados digitados sem salvá-los na configuração da sessão.
            valores = ler_configuracao()
            if valores is not None:
                self._testar_configuracao(valores, janela)

        def aplicar():
            # Guarda os dados válidos apenas durante a sessão atual do aplicativo.
            valores = ler_configuracao()
            if valores is not None:
                self.configuracao = valores
                LOGGER.info(
                    "Configuração PostgreSQL aplicada: servidor=%s porta=%s banco=%s usuário=%s.",
                    valores["host"],
                    valores["port"],
                    valores["dbname"],
                    valores["user"],
                )
                self.status.set(
                    f"Configuração aplicada para {valores['host']}:{valores['port']}/{valores['dbname']} "
                    "nesta sessão."
                )
                janela.destroy()

        botoes = ttk.Frame(formulario)
        botoes.grid(row=len(campos), column=0, columnspan=2, sticky="e", pady=(14, 0))
        ttk.Button(botoes, text="Testar conexão", command=testar).pack(side="left", padx=(0, 8))
        ttk.Button(botoes, text="Cancelar", command=janela.destroy).pack(side="left", padx=(0, 8))
        ttk.Button(botoes, text="Aplicar", command=aplicar).pack(side="left")

    def testar_conexao_salva(self):
        # Testa a configuração atual ou pede que ela seja preenchida primeiro.
        if not self._exigir_configuracao():
            return
        self._testar_configuracao(self.configuracao, self.root)

    def _exigir_configuracao(self):
        # Impede operações no banco até os dados de acesso serem aplicados.
        if self.configuracao is not None:
            return True

        LOGGER.warning("Operação solicitada sem configuração PostgreSQL aplicada.")
        messagebox.showinfo(
            "Configuração necessária",
            "Informe e aplique os dados de acesso ao PostgreSQL para continuar.",
            parent=self.root,
        )
        self.abrir_configuracao()
        return False

    def _testar_configuracao(self, configuracao, parent):
        # Abre uma conexão curta e executa SELECT 1 para validar o acesso.
        LOGGER.info(
            "Teste de conexão iniciado: servidor=%s porta=%s banco=%s usuário=%s.",
            configuracao["host"],
            configuracao["port"],
            configuracao["dbname"],
            configuracao["user"],
        )
        try:
            with psycopg.connect(**configuracao, connect_timeout=5) as conexao:
                with conexao.cursor() as cursor:
                    cursor.execute("SELECT 1")
                    cursor.fetchone()
        except psycopg.Error as erro:
            self.status.set("Falha no teste de conexão.")
            LOGGER.error("Teste de conexão falhou (%s).", type(erro).__name__)
            messagebox.showerror("Falha na conexão", str(erro), parent=parent)
            return False

        LOGGER.info("Teste de conexão concluído com sucesso.")
        self.status.set(
            f"Conexão testada com sucesso: {configuracao['host']}:{configuracao['port']}/"
            f"{configuracao['dbname']}"
        )
        messagebox.showinfo(
            "Conexão OK", "Conexão com o PostgreSQL realizada com sucesso.", parent=parent
        )
        return True

    def abrir_execucao_comandos(self):
        # Cria a tela para executar SELECT, INSERT, UPDATE ou DELETE.
        if not self._exigir_configuracao():
            return

        LOGGER.info("Tela de execução de comandos SQL aberta.")
        janela = tk.Toplevel(self.root)
        janela.title("Executar comandos SQL")
        janela.geometry("900x600")
        janela.minsize(640, 420)
        janela.transient(self.root)

        conteudo = ttk.Frame(janela, padding=12)
        conteudo.pack(fill="both", expand=True)
        ttk.Label(
            conteudo,
            text="Comandos permitidos: SELECT, INSERT, UPDATE e DELETE.",
        ).pack(anchor="w")

        editor = scrolledtext.ScrolledText(conteudo, height=7, wrap="none")
        editor.pack(fill="x", pady=(6, 8))
        editor.insert("1.0", "SELECT * FROM sua_tabela LIMIT 100;")

        barra = ttk.Frame(conteudo)
        barra.pack(fill="x", pady=(0, 8))
        status = tk.StringVar(value="Informe um comando SQL e clique em Executar comando.")
        # Exportação XLS temporariamente desativada:
        # resultado = {"colunas": [], "linhas": []}
        # botao_salvar = ttk.Button(
        #     barra,
        #     text="Salvar resultado XLS...",
        #     command=lambda: self.salvar_resultado_xls(resultado, janela),
        #     state="disabled",
        # )
        # botao_salvar.pack(side="right")
        ttk.Button(
            barra,
            text="Executar comando",
            command=lambda: self.executar_comando_sql(editor, tabela, status, janela),
        ).pack(side="left")
        ttk.Label(barra, textvariable=status).pack(side="left", padx=12)

        tabela = ttk.Treeview(conteudo, show="headings")
        barra_vertical = ttk.Scrollbar(conteudo, orient="vertical", command=tabela.yview)
        barra_horizontal = ttk.Scrollbar(conteudo, orient="horizontal", command=tabela.xview)
        tabela.configure(yscrollcommand=barra_vertical.set, xscrollcommand=barra_horizontal.set)
        tabela.pack(side="top", fill="both", expand=True)
        barra_vertical.pack(side="right", fill="y")
        barra_horizontal.pack(side="bottom", fill="x")

    @staticmethod
    def _identificar_comando_sql(consulta):
        # Ignora comentários iniciais e identifica o primeiro comando SQL.
        restante = consulta.lstrip()
        while restante.startswith("--") or restante.startswith("/*"):
            if restante.startswith("--"):
                fim_comentario = restante.find("\n")
                if fim_comentario == -1:
                    return ""
                restante = restante[fim_comentario + 1 :].lstrip()
            else:
                fim_comentario = restante.find("*/", 2)
                if fim_comentario == -1:
                    return ""
                restante = restante[fim_comentario + 2 :].lstrip()

        correspondencia = re.match(r"[A-Za-z]+", restante)
        return correspondencia.group(0).upper() if correspondencia else ""

    def executar_comando_sql(self, editor, tabela, status, parent):
        # Valida o comando antes de enviar o SQL ao PostgreSQL.
        consulta = editor.get("1.0", "end").strip()
        if not consulta:
            LOGGER.warning("Execução SQL solicitada com o campo vazio.")
            messagebox.showwarning("Comando vazio", "Digite um comando SQL.", parent=parent)
            return
        comando = self._identificar_comando_sql(consulta)
        if comando not in {"SELECT", "INSERT", "UPDATE", "DELETE"}:
            LOGGER.warning("Comando SQL não permitido: tipo=%s.", comando or "desconhecido")
            messagebox.showerror(
                "Comando não permitido",
                "Este menu aceita somente SELECT, INSERT, UPDATE e DELETE.",
                parent=parent,
            )
            return

        # Pede confirmação extra antes de executar comandos que alteram ou removem dados.
        if comando in {"UPDATE", "DELETE"} and not messagebox.askyesno(
            "Confirmar alteração",
            f"O comando {comando} pode alterar ou remover dados. Deseja continuar?",
            parent=parent,
        ):
            LOGGER.info("Execução do comando %s cancelada pelo usuário.", comando)
            return

        tabela.delete(*tabela.get_children())
        nomes_colunas = []
        linhas = []
        registros_afetados = 0

        # Executa dentro de uma transação; o contexto confirma ao terminar sem erro.
        LOGGER.info("Comando SQL iniciado: tipo=%s.", comando)
        try:
            with psycopg.connect(**self.configuracao, connect_timeout=5) as conexao:
                with conexao.cursor() as cursor:
                    # Exemplo de bloco de comando para consulta direta no código:
                    # cursor.execute("SELECT id, nome FROM sua_tabela LIMIT 20")
                    # resultado = cursor.fetchall()

                    cursor.execute(consulta, prepare=True)
                    registros_afetados = max(cursor.rowcount, 0)
                    if cursor.description is None:
                        nomes_colunas = []
                    else:
                        nomes_colunas = [coluna.name for coluna in cursor.description]
                        linhas = cursor.fetchmany(1001)
        except psycopg.Error as erro:
            status.set("Falha ao executar o comando.")
            LOGGER.error("Comando SQL falhou: tipo=%s erro=%s.", comando, type(erro).__name__)
            messagebox.showerror("Erro no comando", str(erro), parent=parent)
            return

        # Confirma operações de escrita somente depois que a transação foi concluída.
        if comando in {"INSERT", "UPDATE", "DELETE"}:
            LOGGER.info(
                "Comando SQL concluído: tipo=%s registros_afetados=%s.",
                comando,
                registros_afetados,
            )
            status.set(
                f"{comando} realizado com sucesso; "
                f"{registros_afetados} registro(s) afetado(s)."
            )
            messagebox.showinfo(
                "Comando realizado com sucesso",
                f"O comando {comando} foi realizado com sucesso.\n"
                f"Registros afetados: {registros_afetados}.",
                parent=parent,
            )

        # Comandos sem retorno de colunas terminam após atualizar o status.
        if not nomes_colunas:
            tabela.configure(columns=())
            return

        # Mostra na grade as colunas e as primeiras 1.000 linhas retornadas.
        tabela.delete(*tabela.get_children())
        identificadores = [f"coluna_{indice}" for indice in range(len(nomes_colunas))]
        tabela.configure(columns=identificadores)
        for identificador, nome in zip(identificadores, nomes_colunas):
            tabela.heading(identificador, text=nome)
            tabela.column(identificador, width=150, minwidth=80, stretch=True)

        limite = 1000
        for linha in linhas[:limite]:
            tabela.insert(
                "",
                "end",
                values=[
                    "" if valor is None else str(self._decodificar_bytes(valor))
                    for valor in linha
                ],
            )

        quantidade = min(len(linhas), limite)
        complemento = " (exibindo as primeiras 1000)" if len(linhas) > limite else ""
        if comando == "SELECT":
            LOGGER.info("Consulta SELECT concluída: linhas=%s.", len(linhas))
            status.set(f"{quantidade} linha(s) retornada(s){complemento}.")

    # Exportação XLS temporariamente desativada:
    # def salvar_resultado_xls(self, resultado, parent):
    #     if not resultado["colunas"]:
    #         messagebox.showwarning(
    #             "Sem resultado", "Execute um SELECT antes de salvar.", parent=parent
    #         )
    #         return
    #
    #     caminho = filedialog.asksaveasfilename(
    #         parent=parent,
    #         title="Salvar resultado da consulta",
    #         initialfile="resultado_consulta.xls",
    #         defaultextension=".xls",
    #         filetypes=[("Planilha Excel 97-2003", "*.xls")],
    #     )
    #     if not caminho:
    #         return
    #
    #     try:
    #         pasta_trabalho = xlwt.Workbook(encoding="utf-8")
    #         planilha = pasta_trabalho.add_sheet("Resultado")
    #         estilo_cabecalho = xlwt.easyxf("font: bold on")
    #
    #         for coluna, nome in enumerate(resultado["colunas"]):
    #             planilha.write(0, coluna, nome, estilo_cabecalho)
    #
    #         for indice_linha, linha in enumerate(resultado["linhas"], start=1):
    #             for indice_coluna, valor in enumerate(linha):
    #                 if valor is None:
    #                     valor = ""
    #                 elif not isinstance(valor, (str, int, float, bool)):
    #                     valor = str(valor)
    #                 planilha.write(indice_linha, indice_coluna, valor)
    #
    #         pasta_trabalho.save(caminho)
    #     except Exception as erro:
    #         messagebox.showerror("Erro ao salvar XLS", str(erro), parent=parent)
    #         return
    #
    #     messagebox.showinfo(
    #         "Arquivo salvo", f"Resultado salvo em:\n{caminho}", parent=parent
    #     )


# Cria a janela principal apenas quando este arquivo é executado diretamente.
if __name__ == "__main__":
    try:
        janela_principal = tk.Tk()
        app = ImportadorApp(janela_principal)
        janela_principal.mainloop()
    except Exception as erro:
        LOGGER.critical("Falha fatal no aplicativo (%s).", type(erro).__name__)
        raise
    else:
        LOGGER.info("Aplicativo encerrado normalmente.")